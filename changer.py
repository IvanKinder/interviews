import openpyxl

wb = openpyxl.load_workbook('DDY.xlsx')

print(wb.sheetnames)
data_sheet = wb['данные за 2 кв.']
classes_sheet = wb['lib класс']
rooms_sheet = wb['lib комнатность']

id_in_data = {}
for i in range(2, data_sheet.max_row):
    if data_sheet.cell(row=i, column=32).value == '':
        continue
    id_in_data[data_sheet.cell(row=i, column=32).value] = (data_sheet.cell(row=i, column=1).value, data_sheet.cell(row=i, column=34).value, data_sheet.cell(row=i, column=11).value, data_sheet.cell(row=i, column=10).value)

id_in_classes = {}
for i in range(2, classes_sheet.max_row):
    if classes_sheet.cell(row=i, column=3).value == '':
        continue
    id_in_classes[classes_sheet.cell(row=i, column=3).value] = (classes_sheet.cell(row=i, column=1).value, data_sheet.cell(row=i, column=4).value)

id_in_rooms = []
for i in range(2, rooms_sheet.max_row):
    id_in_rooms.append([rooms_sheet.cell(row=i, column=1).value, rooms_sheet.cell(row=i, column=2).value, rooms_sheet.cell(row=i, column=3).value, rooms_sheet.cell(row=i, column=4).value])

wb.active = wb['данные за 2 кв.']

data_sheet['BA1'].value = 'Данные о классе'
for i in range(2, data_sheet.max_row):
    if data_sheet.cell(row=i, column=32).value not in id_in_classes.keys():
        data_sheet.cell(row=i, column=53).value = 'Нет такого класса'
    elif id_in_data[data_sheet.cell(row=i, column=32).value][:1] == id_in_classes[data_sheet.cell(row=i, column=32).value]:
        data_sheet.cell(row=i, column=53).value = 'Всё ок'
    else:
        data_sheet.cell(row=i, column=53).value = 'Неверный класс'


for value in id_in_data.values():
    if value[2]:
        print(value)
        for room in id_in_rooms:
            if str(room[1]) == value[2] and room[2] <= value[3] <= room[3]:
                print(room[0])
    else:
        print('Нет комнатности')
print(id_in_rooms)

# wb.save('DDY.xlsx')

